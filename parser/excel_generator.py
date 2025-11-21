from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import logging
from parsers.models import WeatherForecast, DailyWeather
from config import EXCEL_OUTPUT_DIR

logger = logging.getLogger(__name__)

def get_header_style():
    return {
        'font': Font(bold=True, size=14, color="FFFFFF"),
        'fill': PatternFill(start_color="366092", end_color="366092", fill_type="solid"),
        'alignment': Alignment(horizontal="center", vertical="center"),
        'border': Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    }

def get_day_header_style():
    return {
        'font': Font(bold=True, size=12),
        'fill': PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
        'alignment': Alignment(horizontal="center", vertical="center")
    }

def get_table_header_style():
    return {
        'font': Font(bold=True),
        'fill': PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid"),
        'alignment': Alignment(horizontal="center", vertical="center"),
        'border': Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    }

def get_cell_border():
    return Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

def get_warning_style():
    return {
        'font': Font(bold=True, color="FF0000"),
        'fill': PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
        'alignment': Alignment(horizontal="center", vertical="center")
    }

def get_column_widths():
    return {
        'A': 20,
        'B': 18,
        'C': 18,
        'D': 22,
        'E': 15,
        'F': 25,
        'G': 18,
    }

class ExcelGenerator:
    def __init__(self):
        self.workbook = None
        self.worksheet = None
    
    def generate(self, forecast: WeatherForecast) -> Path:
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "Прогноз погоды"
        self._write_header(forecast)
        row = 3
        for day in forecast.days:
            row = self._write_day_data(day, row)
            row += 1
        self._adjust_column_widths()
        filename = f"weather_forecast_{forecast.city}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = EXCEL_OUTPUT_DIR / filename
        self.workbook.save(filepath)
        logger.info(f"Excel файл сохранен: {filepath}")
        return filepath
    
    def _write_header(self, forecast: WeatherForecast):
        header_style = get_header_style()
        self.worksheet.merge_cells('A1:G1')
        cell = self.worksheet['A1']
        cell.value = f"Прогноз погоды на 7 дней для города: {forecast.city}"
        cell.font = header_style['font']
        cell.fill = header_style['fill']
        cell.alignment = header_style['alignment']
        cell.border = header_style['border']
        self.worksheet['A2'] = f"Дата запроса: {forecast.request_date.strftime('%d.%m.%Y %H:%M:%S')}"
        self.worksheet['A2'].font = Font(italic=True)
    
    def _write_day_data(self, day: DailyWeather, start_row: int) -> int:
        current_row = start_row
        day_style = get_day_header_style()
        date_str = day.date.strftime('%d.%m.%Y')
        self.worksheet.merge_cells(f'A{current_row}:G{current_row}')
        cell = self.worksheet[f'A{current_row}']
        cell.value = f"Дата: {date_str} | Магнитное поле: {day.magnetic_field} ({day.magnetic_field_status})"
        cell.font = day_style['font']
        cell.fill = day_style['fill']
        cell.alignment = day_style['alignment']
        current_row += 1
        headers = ['Время суток', 'Температура, °C', 'Ощущаемая, °C', 
                   'Давление, мм рт.ст.', 'Влажность, %', 'Явление', 
                   'Скорость ветра']
        table_header_style = get_table_header_style()
        for col_idx, header in enumerate(headers, start=1):
            cell = self.worksheet.cell(row=current_row, column=col_idx)
            cell.value = header
            cell.font = table_header_style['font']
            cell.fill = table_header_style['fill']
            cell.alignment = table_header_style['alignment']
            cell.border = table_header_style['border']
        current_row += 1
        border = get_cell_border()
        for time_slot in day.time_slots:
            self.worksheet.cell(row=current_row, column=1).value = time_slot.time_of_day
            self.worksheet.cell(row=current_row, column=2).value = time_slot.temperature
            self.worksheet.cell(row=current_row, column=3).value = time_slot.perceived_temperature or time_slot.temperature
            self.worksheet.cell(row=current_row, column=4).value = time_slot.pressure
            self.worksheet.cell(row=current_row, column=5).value = time_slot.humidity
            self.worksheet.cell(row=current_row, column=6).value = time_slot.weather_phenomenon
            self.worksheet.cell(row=current_row, column=7).value = time_slot.wind_speed or "-"
            for col in range(1, 8):
                cell = self.worksheet.cell(row=current_row, column=col)
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")
            current_row += 1
        if day.average_daylight_temperature is not None:
            self.worksheet.merge_cells(f'A{current_row}:B{current_row}')
            cell = self.worksheet[f'A{current_row}']
            cell.value = "Средняя дневная температура:"
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="right", vertical="center")
            self.worksheet.cell(row=current_row, column=3).value = f"{day.average_daylight_temperature} °C"
            self.worksheet.cell(row=current_row, column=3).font = Font(bold=True)
            current_row += 1
        if day.pressure_change_warning:
            warning_style = get_warning_style()
            self.worksheet.merge_cells(f'A{current_row}:G{current_row}')
            cell = self.worksheet[f'A{current_row}']
            cell.value = day.pressure_change_warning
            cell.font = warning_style['font']
            cell.fill = warning_style['fill']
            cell.alignment = warning_style['alignment']
            current_row += 1
        return current_row
    
    def _adjust_column_widths(self):
        column_widths = get_column_widths()
        for col, width in column_widths.items():
            self.worksheet.column_dimensions[col].width = width
    
    def open_file(self, filepath: Path):
        import os
        import platform
        if platform.system() == 'Darwin':
            os.system(f'open "{filepath}"')
        elif platform.system() == 'Windows':
            os.startfile(str(filepath))
        else:
            os.system(f'xdg-open "{filepath}"')
